from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'outputs' / '实习证明.docx'
OUT.parent.mkdir(exist_ok=True)
doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Cm(21), Cm(29.7)
s.top_margin, s.bottom_margin = Cm(3.7), Cm(3)
s.left_margin, s.right_margin = Cm(3.1), Cm(3.1)

def font(style, name, size, bold=False):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0,0,0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),name)

font(doc.styles['Normal'], '宋体', 14)
normal = doc.styles['Normal'].paragraph_format
normal.space_before = Pt(0)
normal.space_after = Pt(0)
normal.line_spacing = Pt(29)
font(doc.styles['Title'], '黑体', 24, True)
t = doc.add_paragraph('实 习 证 明', 'Title')
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
t.paragraph_format.space_after = Pt(38)
t.paragraph_format.line_spacing = Pt(36)

def para(text, indent=True):
    p=doc.add_paragraph(text)
    if indent:
        p.paragraph_format.first_line_indent=Pt(28)
    return p

para('兹证明____________（学校）____________（学院）')
para('_____________________（专业）____________（姓名），', False)
para('自_______年____月____日至_______年____月____日，在我司实习。', False)
para('该同学实习期间工作认真，勤学好问，善于思考，能够举一反三，学以致用。同时，该同学严格遵守我司各项规章制度，与同事和睦相处，得到与其工作的员工对该同学的赞许和肯定。')
p=para('情况属实，特此证明！')
p.paragraph_format.space_after=Pt(82)

for txt in ['单位公章：', '证 明 人：', '联系电话：']:
    p=para(txt, False)
    p.paragraph_format.left_indent=Cm(9)
p=para('日　期：　　年　 月　 日', False)
p.paragraph_format.left_indent=Cm(8.3)

doc.core_properties.title='实习证明'
doc.core_properties.author=''
doc.core_properties.subject='实习证明模板'
for element in doc.styles.element.xpath('.//w:pBdr'):
    element.getparent().remove(element)
for element in doc.element.xpath('.//w:pBdr'):
    element.getparent().remove(element)
doc.save(OUT)
print(OUT)
